#!/usr/bin/env python3
"""Extract and classify an Empower custom-field definition workbook."""

from __future__ import annotations

import argparse
import hashlib
import json
import re
from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook


FUNCTIONS = {
    "ABS", "AVE", "EQ", "EQI", "GT", "GTE", "LT", "LTE", "MAX", "MIN",
    "PROD", "RANGE", "ROUND", "RSD", "SUM",
}

RESERVED = FUNCTIONS | {
    "Area", "Amount", "AS", "B", "Blank", "CCompRef1", "Control", "Dissolved",
    "Error", "Injection", "Label", "Peak", "Retention", "S1", "S2", "SAME",
    "Standard", "Time", "Value", "Width", "SampleWeight", "vt",
}

METADATA_FIELDS = {
    "Accuracy_Level", "Batch_No", "Column Name", "Column Serial Number", "Column_No",
    "Flow_Rate", "Instrument_ID", "LimsSampleId", "LimsTestId", "PostedToLIMS",
    "Product_Name", "ReviewedInLIMS", "Sample Matrix", "SampleName", "Stage", "Level Values",
}

BRIDGE_FIELDS = {
    "Disregard_Factor1", "L1_1", "L2_1", "Label_Claim1", "Label_Claim_1",
    "LOD_Std_1", "LOQ_Percentage1", "MW1_1", "MW2_1", "OVI_Std_1", "RRF1_1",
    "RRF2_1", "Sample_Concentration_PPM_1", "SD1_1", "SD2_1", "SD3_1", "SD4_1",
    "SD5_1", "SD6_1", "SD7_1", "SD_1", "SD_2", "SD_3", "SD_4", "SD_5",
    "SD_6", "SD_7", "Standard_Potency1", "Standard_Potency_1", "Water_Std_1",
    "Weight",
}

INPUT_MEANINGS = {
    "Accuracy_Level": "Sample-level accuracy or recovery level identifier; exact convention is not supplied.",
    "Average_Weight": "Average unit or sample weight used by assay and impurity formulas; units are not supplied.",
    "Batch_No": "Sample batch number.",
    "Column Name": "Default Sample field for chromatographic column name.",
    "Column Serial Number": "Default Sample field for chromatographic column serial number.",
    "Column_No": "Additional sample column identifier; relationship to the two default column fields is not supplied.",
    "Dilution": "Default Sample dilution field, default 1; the listed calculations use Dilution_Factor instead.",
    "Disregard_Factor": "Component-specific multiplier used to derive an impurity disregard threshold.",
    "Dissolution_Media_Volume": "Sample dissolution-media volume used in assay calculations; units are not supplied.",
    "Dose_In_ml": "Sample dose-volume factor used in assay calculations; units are not supplied.",
    "Flow_Rate": "Sample-level flow-rate descriptor stored as text; units and format are not supplied.",
    "Instrument_ID": "Sample instrument identifier.",
    "K": "Sample acceptance-value coefficient, default 2.4.",
    "K_30": "Second acceptance-value input used by the 30-unit branch; exact meaning is not supplied.",
    "L1": "Component acceptance limit, default 15.",
    "L2": "Component individual-unit percentage limit, default 25.",
    "Label_Claim": "Component label claim, default 1; units are not supplied.",
    "Level Values": "Default Sample enumerated level field.",
    "LimsSampleId": "Sample identifier for LIMS integration.",
    "LimsTestId": "Test identifier for LIMS integration.",
    "LOD_Spl": "Sample LOD correction percentage, default 0; the expansion of LOD is not supplied for this assay field.",
    "LOD_Std": "Standard/component LOD correction percentage, default 0; the expansion of LOD is not supplied for this assay field.",
    "LOQ_Percentage": "Component impurity reporting/quantitation threshold percentage, default 0.",
    "MW1": "Component molecular-weight numerator factor, default 1; chemical direction and units remain unconfirmed.",
    "MW2": "Component molecular-weight denominator factor, default 1; chemical direction and units remain unconfirmed.",
    "No_Of_Units": "Number of dosage units represented by the sample, default 1.",
    "OVI_Spl": "Sample OVI correction percentage, default 0; the expansion of OVI is not supplied.",
    "OVI_Std": "Standard/component OVI correction percentage, default 0; the expansion of OVI is not supplied.",
    "Product_Name": "Sample product name.",
    "RRF1": "Component relative-response factor used as the numerator, default 1.",
    "RRF2": "Component relative-response factor used as the denominator, default 1.",
    "Sample Matrix": "Default Sample enumerated matrix field.",
    "Sample_Concentration_PPM": "Component sample concentration in ppm, default 1.",
    "Sample_Potency": "Sample potency percentage used in added-amount calculations; default is blank.",
    "SampleName": "Default-style sample name field stored under the workbook spelling `SampleName`.",
    "SD1": "First Component standard-dilution input, default 1; flask/aliquot role is unconfirmed.",
    "SD2": "Second Component standard-dilution input, default 1; flask/aliquot role is unconfirmed.",
    "SD3": "Third Component standard-dilution input, default 1; flask/aliquot role is unconfirmed.",
    "SD4": "Fourth Component standard-dilution input, default 1; flask/aliquot role is unconfirmed.",
    "SD5": "Fifth Component standard-dilution input, default 1; flask/aliquot role is unconfirmed.",
    "SD6": "Sixth Component standard-dilution input, default 1; flask/aliquot role is unconfirmed.",
    "SD7": "Seventh Component standard-dilution input, default 1; flask/aliquot role is unconfirmed.",
    "Stage": "Sample stage descriptor; likely a calculation or testing stage, but the convention is not supplied.",
    "Standard_Potency": "Component standard potency percentage, default 1 in this workbook.",
    "TD1": "First required Sample dilution input, default 1; flask/aliquot role is unconfirmed.",
    "TD2": "Second required Sample dilution input, default 1; flask/aliquot role is unconfirmed.",
    "TD3": "Third required Sample dilution input, default 1; flask/aliquot role is unconfirmed.",
    "TD4": "Fourth required Sample dilution input, default 1; flask/aliquot role is unconfirmed.",
    "TD5": "Fifth required Sample dilution input, default 1; flask/aliquot role is unconfirmed.",
    "TD6": "Sixth required Sample dilution input, default 1; flask/aliquot role is unconfirmed.",
    "TD7": "Seventh required Sample dilution input, default 1; flask/aliquot role is unconfirmed.",
    "Test_Target_Conc_PPM": "Sample test target concentration in ppm; no formula in this workbook references it.",
    "Water_Spl": "Sample water correction percentage, default 0.",
    "Water_Std": "Standard/component water correction percentage, default 0.",
    "Weight_Per_mL": "Sample weight-per-mL factor used in assay calculations, default 1; units are not supplied.",
}


def serializable(value):
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return value


def clean_formula(value) -> str:
    return value.strip() if isinstance(value, str) else ""


def family_for(name: str) -> str:
    if name in METADATA_FIELDS:
        return "Sample metadata and LIMS"
    if name in BRIDGE_FIELDS or re.fullmatch(r"SD\d|TD\d", name):
        return "Inputs and bridge fields"
    if name in {"Dilution", "Dilution_Factor", "Average_Weight", "Dissolution_Media_Volume", "Dose_In_ml", "No_Of_Units", "Weight_Per_mL", "SampleWeight", "Label_Claim", "LOD_Spl", "LOD_Std", "LOQ_Percentage", "MW1", "MW2", "OVI_Spl", "OVI_Std", "RRF1", "RRF2", "Standard_Potency", "Water_Spl", "Water_Std"}:
        return "Inputs and bridge fields"
    if name.startswith(("LOD_Determ", "LOD_DETERMIN", "LOQ_Determ", "LOQ_DETERMIN")):
        return "LOD and LOQ determination"
    if name in {"AV", "AV_30", "Average_Percentage_Assay", "Higher_Limit_Individual_Unit", "K", "KS", "KS_30", "K_30", "L1", "L1_1", "L2", "L2_1", "Lower_Limit_Individual_Unit", "MAX_ASSAY_SINGLE_AREA", "MIN_ASSAY_SINGLE_AREA", "M_Value", "RSD_Percentage_Assay", "Standard_Deviation_Perc_Assay"}:
        return "Content uniformity and acceptance value"
    if name.startswith(("Added_Amount", "Found_Amount", "Found_Amoun", "Recovery", "AVE_RECOVERY", "AVG_FOUND", "Percentage_Recovery", "Sample_Concentration", "Sample_Potency", "Sample_Response_Mean")) or name == "Test_Target_Conc_PPM":
        return "Accuracy and recovery"
    if name.startswith(("Assay_", "mg_Per_Unit", "Avg_Assay", "AVG_Percentage_Assay")) or name in {"Amount_Label", "Sample_Average_Area", "STD_AVG_Area"}:
        return "Assay and content"
    if "Disregard" in name or "Area_Dis" in name:
        return "Impurity disregard threshold"
    if "LOQ" in name and ("Imp" in name or name.startswith("Total_")):
        return "Impurity LOQ filtering and reporting"
    if "AMV" in name:
        return "Impurity AMV"
    if "Corrected_Area_AN" in name or "Area_Corrected_AN" in name or "Imp_AN" in name or "IMP_AN" in name:
        return "Impurity area normalization"
    if "Impurity" in name or "Impurities" in name or "IMP_Label" in name or name.startswith("Single_Max_Unknown"):
        return "Impurity amount based"
    if name in {"Blank_Response", "Area_Correction", "Response_Correction"}:
        return "Blank and response correction"
    if name.startswith("Standard1_") or name.startswith("Standard2_") or name.startswith("Similarity") or name in {"RESPONSE_RATIO", "PERCENT_STANDARD_RSD", "CUMMULATIVE_PERCENT_RSD", "Standard_Average_Area", "Standard_Average_Area_Cal", "Standard_Weight", "STD_CON"}:
        return "Standards, response, and similarity"
    if name in {"Plate_Count_USP", "Relative_Resolution_HH"}:
        return "Chromatography performance"
    if name == "Percentage_Drug_Released":
        return "Dissolution"
    if name == "Residual_Solvent":
        return "Residual solvent"
    if name in {"Value_Amount_Table", "X_Axis_ppm"}:
        return "Calibration and display helpers"
    return "Unclassified"


def role_for(record: dict) -> str:
    if record["name"] in BRIDGE_FIELDS:
        return "Bridge/copy"
    if record["source"] == "Keyboard":
        return "Manual input"
    if record["source"] == "External":
        return "External input"
    if record["intersample"]:
        return "Intersample calculated"
    if record["field_type"] == "Result":
        return "Result aggregate"
    return "Calculated"


def inferred_meaning(record: dict) -> str:
    name = record["name"]
    if name in INPUT_MEANINGS:
        return INPUT_MEANINGS[name]
    if name in BRIDGE_FIELDS and record["dependencies"]:
        return f"Peak-level copy of `{record['dependencies'][0]}`, likely used for display, reporting, or downstream Peak formulas."
    if record["source"] == "External":
        return "Result-level status supplied by an external LIMS integration."
    readable = name.replace("_", " ")
    if record["intersample"]:
        return f"Cross-sample or cross-injection calculation for {readable}; selector meaning depends on the sample-set labels and processing order."
    if record["field_type"] == "Result" and clean_formula(record["formula"]).startswith(("SUM(", "MAX(")):
        return f"Result-level aggregate for {readable}."
    return f"Calculated field for {readable}; purpose is inferred from its name and exact formula."


def find_dependencies(formula: str, names: set[str], lower_names: dict[str, str]):
    without_strings = re.sub(r'"[^"]*"', " ", formula)
    tokens = re.findall(r"[A-Za-z_][A-Za-z0-9_]*", without_strings)
    dependencies = []
    case_variants = []
    unknown = []
    for token in tokens:
        if token in names:
            if token not in dependencies:
                dependencies.append(token)
            continue
        canonical = lower_names.get(token.lower())
        if canonical:
            if canonical not in dependencies:
                dependencies.append(canonical)
            case_variants.append({"formula_token": token, "defined_name": canonical})
            continue
        if token not in RESERVED and token not in unknown:
            unknown.append(token)
    return dependencies, case_variants, unknown


def is_intersample(formula: str) -> bool:
    if re.search(r"\b(?:SAME|S1|S2|Blank|AS)\.", formula):
        return True
    return bool(re.search(r"(?<!\d)\.%(?:\.|\d)", formula))


def intersample_details(formula: str) -> tuple[str, list[str]]:
    if not is_intersample(formula):
        return "", []
    kind = "summary" if re.search(r"\.(?:AVE|MAX|MIN|SUM|%RSD|PROD)\(", formula) else "non-summary"
    selectors = []
    for match in re.finditer(r"(?<![A-Za-z0-9_])(S1%|SAME|S1|S2|Blank|AS)\.", formula):
        label = match.group(1)
        if label not in selectors:
            selectors.append(label)
    if formula.lstrip().startswith(".%") and "current/omitted label" not in selectors:
        selectors.append("current/omitted label")
    return kind, selectors


def find_cycles(records: list[dict]) -> list[list[str]]:
    graph = {record["name"]: record["dependencies"] for record in records}
    visiting = set()
    visited = set()
    stack = []
    cycles = []

    def visit(node: str) -> None:
        if node in visiting:
            start = stack.index(node)
            cycle = stack[start:] + [node]
            if cycle not in cycles:
                cycles.append(cycle)
            return
        if node in visited:
            return
        visiting.add(node)
        stack.append(node)
        for dependency in graph.get(node, []):
            visit(dependency)
        stack.pop()
        visiting.remove(node)
        visited.add(node)

    for field_name in graph:
        visit(field_name)
    return cycles


def md_cell(value) -> str:
    if value is None or value == "":
        return ""
    return str(value).replace("|", "\\|").replace("\n", "<br>")


def code_cell(value) -> str:
    if value is None or value == "":
        return ""
    return f"`{str(value).strip()}`"


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("xlsx", type=Path)
    parser.add_argument("transcript", type=Path)
    parser.add_argument("catalog", type=Path)
    parser.add_argument("inventory", type=Path)
    args = parser.parse_args()

    source_bytes = args.xlsx.read_bytes()
    source_hash = hashlib.sha256(source_bytes).hexdigest()
    workbook = load_workbook(args.xlsx, data_only=False, read_only=False)
    sheet = workbook[workbook.sheetnames[0]]
    headers = [cell.value for cell in sheet[1]]
    records = []
    for row in sheet.iter_rows(min_row=2, max_col=len(headers), values_only=True):
        if not any(value is not None for value in row):
            continue
        raw = dict(zip(headers, row))
        records.append(
            {
                "number": raw.get(None),
                "name": raw.get("Name"),
                "field_type": raw.get("Field Type"),
                "data_type": raw.get("Type"),
                "width": raw.get("Width"),
                "precision": raw.get("Precision"),
                "minimum": raw.get("Minimum"),
                "maximum": raw.get("Maximum"),
                "default": raw.get("Default"),
                "source": raw.get("Source"),
                "required": bool(raw.get("Required")),
                "formula": clean_formula(raw.get("Formula")),
            }
        )

    names = {record["name"] for record in records}
    lower_names = {name.lower(): name for name in names}
    for record in records:
        deps, variants, unknown = find_dependencies(record["formula"], names, lower_names)
        record["dependencies"] = deps
        record["case_variant_dependencies"] = variants
        record["unresolved_identifiers"] = unknown
        record["intersample"] = is_intersample(record["formula"])
        record["intersample_kind"], record["intersample_selectors"] = intersample_details(record["formula"])
        record["family"] = family_for(record["name"])
        record["role"] = role_for(record)
        record["meaning"] = inferred_meaning(record)

    referenced_by = defaultdict(list)
    for record in records:
        for dependency in record["dependencies"]:
            referenced_by[dependency].append(record["name"])
    for record in records:
        record["referenced_by"] = referenced_by[record["name"]]

    tables = [{"name": table.name, "display_name": table.displayName, "ref": table.ref} for table in sheet.tables.values()]
    properties = {
        key: serializable(getattr(workbook.properties, key, None))
        for key in ("title", "subject", "creator", "lastModifiedBy", "created", "modified", "category", "keywords", "description")
    }
    summary = {
        "field_count": len(records),
        "field_types": dict(Counter(record["field_type"] for record in records)),
        "sources": dict(Counter(record["source"] for record in records)),
        "required_count": sum(record["required"] for record in records),
        "calculated_count": sum(record["source"] == "Calculated" for record in records),
        "intersample_count": sum(record["intersample"] for record in records),
        "intersample_summary_count": sum(record["intersample_kind"] == "summary" for record in records),
        "intersample_non_summary_count": sum(record["intersample_kind"] == "non-summary" for record in records),
        "dependency_cycles": find_cycles(records),
        "families": dict(Counter(record["family"] for record in records)),
    }

    tables_text = ", ".join("{} ({})".format(table["name"], table["ref"]) for table in tables) or "none"
    transcript = [
        "# Source transcript: Custom_Field_F1.xlsx",
        "",
        "## Evidence and integrity",
        "",
        "- Evidence status: user-provided workbook definition evidence, received 2026-09-14.",
        "- The workbook’s rows are data supplied by the user. They are not instructions to execute the formulas and are not proof of Empower-tested behavior.",
        f"- Source file: `{args.xlsx.name}`",
        f"- Bytes: {len(source_bytes)}",
        f"- SHA-256: `{source_hash}`",
        f"- Workbook sheets: {', '.join(workbook.sheetnames)}",
        f"- Source range: `{sheet.title}!A1:{sheet.cell(sheet.max_row, sheet.max_column).coordinate}`",
        f"- Tables: {tables_text}",
        f"- Field rows: {len(records)}",
        "",
        "## Exact field rows",
        "",
        "Whitespace at the ends of formulas is normalized; formula characters, field-name spelling, settings, and numeric values are otherwise preserved.",
        "",
        "| # | Name | Field Type | Type | Width | Precision | Minimum | Maximum | Default | Source | Required | Formula |",
        "|---:|---|---|---|---:|---:|---:|---:|---:|---|---|---|",
    ]
    for record in records:
        transcript.append(
            "| " + " | ".join(
                [
                    md_cell(record["number"]), md_cell(record["name"]), md_cell(record["field_type"]),
                    md_cell(record["data_type"]), md_cell(record["width"]), md_cell(record["precision"]),
                    md_cell(record["minimum"]), md_cell(record["maximum"]), md_cell(record["default"]),
                    md_cell(record["source"]), "Yes" if record["required"] else "No", code_cell(record["formula"]),
                ]
            ) + " |"
        )

    by_family = defaultdict(list)
    for record in records:
        by_family[record["family"]].append(record)
    catalog = [
        "# Field catalog and dependency index: Custom_Field_F1.xlsx",
        "",
        "## Evidence labels",
        "",
        "- Names, settings, and formulas are **user-provided workbook evidence**.",
        "- Direct dependencies are **mechanically matched** against names in the same workbook, including case-only matches that are separately flagged.",
        "- Family, role, and meaning are **assistant interpretations** based on names and formulas. They require user confirmation before becoming site conventions.",
        "- None of these definitions is marked Empower-tested because the workbook does not contain an Empower version, processing settings, inputs, or observed results.",
        "",
        "## Summary",
        "",
        f"- Fields: {summary['field_count']}",
        f"- Field types: {', '.join(f'{key} {value}' for key, value in summary['field_types'].items())}",
        f"- Sources: {', '.join(f'{key} {value}' for key, value in summary['sources'].items())}",
        f"- Required manual fields: {summary['required_count']}",
        f"- Intersample or inter-injection formulas detected: {summary['intersample_count']}",
        "",
    ]
    for family in sorted(by_family):
        catalog.extend(
            [
                f"## {family}",
                "",
                "| # | Field | Type/source | Role | Direct custom-field dependencies | Intersample | Inferred meaning | Formula |",
                "|---:|---|---|---|---|---|---|---|",
            ]
        )
        for record in by_family[family]:
            dependencies = ", ".join(f"`{name}`" for name in record["dependencies"])
            if record["case_variant_dependencies"]:
                variants = ", ".join(
                    f"`{item['formula_token']}`→`{item['defined_name']}`"
                    for item in record["case_variant_dependencies"]
                )
                dependencies = f"{dependencies}; case-only: {variants}" if dependencies else f"case-only: {variants}"
            catalog.append(
                "| " + " | ".join(
                    [
                        md_cell(record["number"]), code_cell(record["name"]),
                        md_cell(f"{record['field_type']} / {record['source']}"), md_cell(record["role"]),
                        dependencies, "Yes" if record["intersample"] else "No",
                        md_cell(record["meaning"]), code_cell(record["formula"]),
                    ]
                ) + " |"
            )
        catalog.append("")

    inventory = {
        "source": args.xlsx.name,
        "source_bytes": len(source_bytes),
        "source_sha256": source_hash,
        "properties": properties,
        "sheets": workbook.sheetnames,
        "active_sheet": sheet.title,
        "source_range": f"{sheet.title}!A1:{sheet.cell(sheet.max_row, sheet.max_column).coordinate}",
        "tables": tables,
        "summary": summary,
        "fields": records,
    }

    for output in (args.transcript, args.catalog, args.inventory):
        output.parent.mkdir(parents=True, exist_ok=True)
    args.transcript.write_text("\n".join(transcript) + "\n", encoding="utf-8")
    args.catalog.write_text("\n".join(catalog), encoding="utf-8")
    args.inventory.write_text(json.dumps(inventory, indent=2, ensure_ascii=False), encoding="utf-8")


if __name__ == "__main__":
    main()
